"""Integration tests for spreadsheet formula-injection protection."""

from pathlib import Path

from openpyxl import load_workbook

from odysafe_threatmap.exporters.excel.tables import ExcelColumn, create_excel_table
from odysafe_threatmap.exporters.excel.workbook import ExcelWorkbook


def test_dangerous_text_is_not_written_as_excel_formula(tmp_path: Path) -> None:
    """Workbook output preserves formula-like external text as literal strings."""
    path = tmp_path / "security.xlsx"
    with ExcelWorkbook(path) as workbook:
        sheet = workbook.add_sheet("Data")
        create_excel_table(
            workbook,
            sheet,
            [{"value": "=SUM(A1:A10)"}, {"value": "+A1"}, {"value": "-A1"}, {"value": "@A1"}],
            0,
            0,
            [ExcelColumn("value", "Value")],
        )
    worksheet = load_workbook(path, data_only=False)["Data"]
    assert all(worksheet.cell(row, 1).data_type != "f" for row in range(2, 6))
