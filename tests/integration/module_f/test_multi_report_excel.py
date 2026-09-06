"""Offline integration test for multi-report Excel output."""

from pathlib import Path

from openpyxl import load_workbook

from odysafe_threatmap.application.multi_report import AggregateOptions, MultiReportService
from odysafe_threatmap.exporters.excel.multi_report import build_aggregate_workbook
from odysafe_threatmap.infrastructure.attack.repository import AttackRepositoryImpl
from odysafe_threatmap.infrastructure.extraction.iocsearcher_adapter import IocsearcherAdapter
from odysafe_threatmap.infrastructure.extraction.pipeline import ExtractionPipeline


class NoSupplementalExtraction:
    """Offline supplemental extractor test double."""

    def extract(self, file_path: Path):
        """Return no supplemental extraction results."""
        from odysafe_threatmap.domain.models import ExtractionResult

        return ExtractionResult()


def test_aggregate_workbook_has_required_sheets(tmp_path: Path) -> None:
    """Multiple reports generate an inspectable aggregate workbook."""
    root = Path(__file__).parents[3]
    repository = AttackRepositoryImpl(root / "tests/fixtures/attack/enterprise-attack-test.json")
    service = MultiReportService(
        ExtractionPipeline(IocsearcherAdapter(), NoSupplementalExtraction(), repository), repository
    )
    result = service.aggregate_reports(
        [root / "tests/fixtures/reports/report_test_1.txt", root / "tests/fixtures/reports/report_test_2.txt"],
        AggregateOptions(),
    )
    path = tmp_path / "aggregate.xlsx"
    build_aggregate_workbook(result, path, "test")
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Synthèse", "Rapports", "IOCs", "TTPs", "Fausse corroboration", "_Meta"]
    assert workbook["_Meta"].sheet_state == "hidden"
    assert workbook["Rapports"].tables
    meta_values = {
        row[0].value: row[1].value
        for row in workbook["_Meta"].iter_rows(min_row=2, max_col=2)
        if row[0].value is not None
    }
    assert meta_values["input_hashes"] != "N/A"
    for report in result.reports:
        assert meta_values[f"input_sha256:{report.filename}"] == report.sha256
