"""Offline end-to-end tests for Module B Excel output."""

from pathlib import Path

from openpyxl import load_workbook

from odysafe_threatmap.application.actor_snapshot import ActorSnapshotOptions, ActorSnapshotService
from odysafe_threatmap.cli.commands.actor import _output_label
from odysafe_threatmap.exporters.excel.actor_snapshot import build_actor_workbook
from odysafe_threatmap.infrastructure.attack.repository import AttackRepositoryImpl


def test_multi_actor_snapshot_has_required_sheets(tmp_path: Path) -> None:
    """Two actors produce a readable workbook including the comparison matrix."""
    root = Path(__file__).parents[3]
    service = ActorSnapshotService(AttackRepositoryImpl(root / "tests/fixtures/attack/enterprise-attack-test.json"))
    result = service.build_snapshot(["G1001", "G1002"], ActorSnapshotOptions(include_campaigns=True))
    output_path = tmp_path / "actors.xlsx"

    build_actor_workbook(result, output_path, "odysafe-threatmap actor snapshot")

    workbook = load_workbook(output_path)
    assert workbook.sheetnames == [
        "Synthèse",
        "Profil",
        "TTPs",
        "Logiciels - Malwares",
        "Campagnes",
        "Heatmap tactiques",
        "Comparaison acteurs",
        "_Meta",
    ]
    assert workbook["_Meta"].sheet_state == "hidden"
    meta_values = {
        row[0].value: row[1].value
        for row in workbook["_Meta"].iter_rows(min_row=2, max_col=2)
        if row[0].value is not None
    }
    assert meta_values["attack_stix_version"] == "2.0"
    assert meta_values["input_hashes"].startswith("Not applicable")
    assert meta_values["mitreattack_python_version"] != "N/A"
    assert workbook["Profil"].tables
    assert workbook["TTPs"].tables
    assert workbook["Synthèse"]["A3"].value == "Actors"
    assert workbook["Synthèse"].column_dimensions["U"].hidden is not True
    selected_actor_counts = [workbook["Synthèse"].cell(row=row, column=2).value for row in range(112, 122)]
    assert max(value for value in selected_actor_counts if isinstance(value, int)) <= len(result.actors)
    assert workbook["Profil"].freeze_panes == "A2"
    assert workbook["TTPs"].freeze_panes == "A2"
    comparison = workbook["Comparaison acteurs"]
    assert comparison["A1"].value == "Actor Technique Comparison"
    assert comparison.freeze_panes == "A10"
    assert comparison.sheet_view.zoomScale == 85
    assert len(comparison._charts) == 1
    assert _output_label(result).startswith("comparison_2-groups_")
    assert result.actors[0].name not in _output_label(result)
    assert (tmp_path / "manifest.json").is_file()


def test_single_actor_snapshot_omits_comparison_sheet(tmp_path: Path) -> None:
    """A single actor does not produce an unnecessary comparison matrix."""
    root = Path(__file__).parents[3]
    service = ActorSnapshotService(AttackRepositoryImpl(root / "tests/fixtures/attack/enterprise-attack-test.json"))
    result = service.build_snapshot(["G1001"], ActorSnapshotOptions())
    output_path = tmp_path / "actor.xlsx"

    build_actor_workbook(result, output_path, "odysafe-threatmap actor snapshot")

    assert "Comparaison acteurs" not in load_workbook(output_path).sheetnames
