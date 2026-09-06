"""Structural and presentation-quality tests for Module A workbooks."""

from pathlib import Path

from openpyxl import load_workbook

from odysafe_threatmap.application.report_bundle import ReportBundleService, ReportMetadataOptions
from odysafe_threatmap.exporters.excel.report_bundle import write_report_bundle_workbook
from odysafe_threatmap.infrastructure.attack.repository import AttackRepositoryImpl
from odysafe_threatmap.infrastructure.extraction.iocsearcher_adapter import IocsearcherAdapter
from odysafe_threatmap.infrastructure.extraction.pipeline import ExtractionPipeline


class NoSupplementalExtraction:
    """Offline test double for supplemental extraction."""

    def extract(self, file_path: Path):
        """Return an empty supplemental extraction result."""
        from odysafe_threatmap.domain.models import ExtractionResult

        return ExtractionResult()


def test_module_a_text_and_numeric_output_quality(tmp_path: Path) -> None:
    """Visible descriptions are clean and overlap percentages remain numeric."""
    root = Path(__file__).parents[3]
    repository = AttackRepositoryImpl(root / "tests/fixtures/attack/enterprise-attack-test.json")
    service = ReportBundleService(
        ExtractionPipeline(IocsearcherAdapter(), NoSupplementalExtraction(), repository), repository
    )
    result = service.build_bundle(root / "tests/fixtures/reports/report_test_1.txt", ReportMetadataOptions())
    path = tmp_path / "quality.xlsx"
    write_report_bundle_workbook(path, result, "test")
    workbook = load_workbook(path)
    assert workbook["_Meta"].sheet_state == "hidden"
    assert workbook["_Meta"]["B12"].value not in {None, "N/A"}
    assert workbook["TTPs"].freeze_panes is not None
    for row in workbook["TTPs"].iter_rows(values_only=True):
        for value in row:
            if isinstance(value, str):
                assert "<code>" not in value
                assert "(Citation:" not in value
    overlap = workbook["Groupes associés"]["E5"].value
    assert isinstance(overlap, (float, int))
