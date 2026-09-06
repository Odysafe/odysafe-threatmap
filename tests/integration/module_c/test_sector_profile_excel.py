"""Offline end-to-end tests for Sector Threat Profile Excel output."""

from pathlib import Path

from openpyxl import load_workbook

from odysafe_threatmap.application.sector_profile import SectorProfileOptions, SectorProfileService
from odysafe_threatmap.config.loader import load_actor_metadata
from odysafe_threatmap.exporters.excel.sector_profile import build_sector_workbook
from odysafe_threatmap.infrastructure.attack.repository import AttackRepositoryImpl


def test_sector_profile_workbook_has_required_sheets(tmp_path: Path) -> None:
    """A locally mapped sector creates the required readable workbook."""
    root = Path(__file__).parents[3]
    service = SectorProfileService(
        AttackRepositoryImpl(root / "tests/fixtures/attack/enterprise-attack-test.json"),
        load_actor_metadata(root / "src/odysafe_threatmap/config/defaults/actor_metadata.yaml"),
    )
    result = service.build_profile(
        ["financial"],
        SectorProfileOptions(
            include_regions=True,
            tactic_weights={"Initial Access": 4},
            priority_thresholds={"critical_threshold": 150, "high_threshold": 100, "moderate_threshold": 50},
        ),
    )
    path = tmp_path / "sector.xlsx"
    build_sector_workbook(result, path, "test")
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Synthèse", "Secteur", "Acteurs", "TTPs", "Matrice de risque", "Régions", "_Meta"]
    assert workbook["_Meta"].sheet_state == "hidden"
    assert workbook["Matrice de risque"]["A1"].value == "Odysafe local prioritization — configurable heuristic"
    assert workbook["TTPs"].tables
