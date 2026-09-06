"""End-to-end offline tests for Module A Excel output."""

from pathlib import Path

from openpyxl import load_workbook

from odysafe_threatmap.application.report_bundle import ReportBundleService, ReportMetadataOptions
from odysafe_threatmap.exporters.excel.report_bundle import write_report_bundle_workbook
from odysafe_threatmap.infrastructure.attack.repository import AttackRepositoryImpl
from odysafe_threatmap.infrastructure.extraction.iocsearcher_adapter import IocsearcherAdapter
from odysafe_threatmap.infrastructure.extraction.pipeline import ExtractionPipeline


class NoSupplementalExtraction:
    """Offline test double for supplemental extraction."""

    def extract(self, file_path: Path):
        """Return no supplemental results."""
        from odysafe_threatmap.domain.models import ExtractionResult

        return ExtractionResult()


def test_report_bundle_excel_has_required_sheet_order_and_structure(tmp_path: Path) -> None:
    """A report produces a readable workbook with all Module A sheets."""
    root = Path(__file__).parents[3]
    repository = AttackRepositoryImpl(root / "tests" / "fixtures" / "attack" / "enterprise-attack-test.json")
    service = ReportBundleService(
        ExtractionPipeline(IocsearcherAdapter(), NoSupplementalExtraction(), repository), repository
    )
    result = service.build_bundle(
        root / "tests" / "fixtures" / "reports" / "report_test_1.txt", ReportMetadataOptions(name="Test report")
    )
    output_path = tmp_path / "report.xlsx"

    write_report_bundle_workbook(output_path, result, "odysafe-threatmap report build")

    workbook = load_workbook(output_path)
    assert workbook.sheetnames == [
        "Synthèse",
        "Rapport",
        "Résumé IOCs",
        "IOCs",
        "TTPs",
        "Couverture Tactiques",
        "Détection",
        "Mitigations",
        "Groupes associés",
        "Acteurs - Campagnes",
        "_Meta",
    ]
    assert workbook["_Meta"].sheet_state == "hidden"
    assert workbook["Rapport"].tables
    assert workbook["IOCs"].tables
    assert str(workbook["Synthèse"]["A6"].value).startswith("IOCs uniques\n")
    assert workbook["Synthèse"]["K7"].value == "CTI Observation Density"
    assert "observations / 1,000 words" in workbook["Synthèse"]["K8"].value
    assert "concentration indicator only" in workbook["Synthèse"]["A10"].value
    report_values = {
        row[0].value: row[1].value
        for row in workbook["Rapport"].iter_rows(min_row=2, max_col=2)
        if row[0].value is not None
    }
    assert report_values["CTI Observation Density"].endswith(" observations / 1,000 words")
    assert report_values["Counting basis"] == "Occurrences, not unique values"
    meta_values = {
        row[0].value: row[1].value
        for row in workbook["_Meta"].iter_rows(min_row=2, max_col=2)
        if row[0].value is not None
    }
    assert meta_values["cti_density_formula"].startswith("(IOC occurrences + explicit TTP occurrences)")
    assert meta_values["attack_stix_version"] == "2.0"
    assert meta_values["attack_source_type"] == "Local MITRE CTI Enterprise ATT&CK STIX bundle"
    assert meta_values["attack_source_ref"] == meta_values["attack_bundle_path"]
    assert meta_values["input_hashes"] != "N/A"
    assert meta_values[f"input_sha256:{result.report_metadata.filename}"] == result.report_metadata.sha256
    assert "ne constitue pas une attribution" in workbook["Groupes associés"]["A1"].value
    assert (tmp_path / "manifest.json").is_file()


def test_poor_report_still_produces_all_sheets(tmp_path: Path) -> None:
    """A poor report is exported successfully with empty-data banners."""
    root = Path(__file__).parents[3]
    repository = AttackRepositoryImpl(root / "tests" / "fixtures" / "attack" / "enterprise-attack-test.json")
    service = ReportBundleService(
        ExtractionPipeline(IocsearcherAdapter(), NoSupplementalExtraction(), repository), repository
    )
    result = service.build_bundle(root / "tests" / "fixtures" / "reports" / "report_poor.txt", ReportMetadataOptions())
    output_path = tmp_path / "poor.xlsx"

    write_report_bundle_workbook(output_path, result, "odysafe-threatmap report build")

    workbook = load_workbook(output_path)
    assert workbook["TTPs"]["A1"].value == "ℹ Aucune donnée disponible."
    assert workbook["IOCs"]["A1"].value == "ℹ Aucune donnée disponible."
