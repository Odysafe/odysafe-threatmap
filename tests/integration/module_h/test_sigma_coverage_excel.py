"""Offline integration test for Sigma coverage Excel output."""

from pathlib import Path

from openpyxl import load_workbook

from odysafe_threatmap.application.sigma_coverage import SigmaCoverageOptions, SigmaCoverageService
from odysafe_threatmap.exporters.excel.sigma_coverage import build_coverage_workbook
from odysafe_threatmap.infrastructure.attack.repository import AttackRepositoryImpl
from odysafe_threatmap.infrastructure.extraction.iocsearcher_adapter import IocsearcherAdapter
from odysafe_threatmap.infrastructure.extraction.pipeline import ExtractionPipeline


class NoSupplementalExtraction:
    """Offline supplemental extractor test double."""

    def extract(self, file_path: Path):
        """Return no supplemental extraction results."""
        from odysafe_threatmap.domain.models import ExtractionResult

        return ExtractionResult()


def test_sigma_coverage_workbook_has_required_sheets(tmp_path: Path) -> None:
    """Strict tag coverage produces a readable workbook."""
    root = Path(__file__).parents[3]
    repository = AttackRepositoryImpl(root / "tests/fixtures/attack/enterprise-attack-test.json")
    service = SigmaCoverageService(
        ExtractionPipeline(IocsearcherAdapter(), NoSupplementalExtraction(), repository),
        repository,
        {"Initial Access": 4},
    )
    result = service.analyze_coverage(
        root / "tests/fixtures/reports/report_test_1.txt",
        root / "tests/fixtures/sigma",
        SigmaCoverageOptions(include_rules=True),
    )
    path = tmp_path / "coverage.xlsx"
    build_coverage_workbook(result, path, "test")
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Synthèse", "Couverture", "Règles Sigma", "_Meta"]
    assert workbook["_Meta"].sheet_state == "hidden"
    assert workbook["Couverture"].tables
    meta_values = {
        row[0].value: row[1].value
        for row in workbook["_Meta"].iter_rows(min_row=2, max_col=2)
        if row[0].value is not None
    }
    assert meta_values["input_hashes"] != "N/A"
    assert len([key for key in meta_values if str(key).startswith("input_sha256:")]) == len(result.input_hashes)
