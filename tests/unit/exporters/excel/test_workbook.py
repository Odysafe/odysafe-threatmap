"""Openpyxl-backed tests for the reusable Excel framework."""

import json
from pathlib import Path

from openpyxl import load_workbook

from odysafe_threatmap.exporters.excel.charts import create_bar_chart
from odysafe_threatmap.exporters.excel.sheets.meta_sheet import write_meta_sheet
from odysafe_threatmap.exporters.excel.sheets.summary_sheet import write_summary_sheet
from odysafe_threatmap.exporters.excel.tables import ExcelColumn, create_excel_table
from odysafe_threatmap.exporters.excel.workbook import ExcelWorkbook


def test_workbook_writes_tables_kpis_meta_and_manifest(tmp_path: Path) -> None:
    """Generated workbooks are readable and contain shared framework artifacts."""
    filename = tmp_path / "framework.xlsx"
    metadata = {
        "inputs": ["report.txt"],
        "hashes": {"report.txt": "abc"},
        "versions": {"odysafe-threatmap": "0.1.0"},
        "attack_dataset": {"sha256": "def"},
        "command": "odysafe-threatmap report build report.txt",
    }
    with ExcelWorkbook(filename, metadata=metadata) as workbook:
        report = workbook.add_sheet("Report")
        workbook.write_title(report, 0, 0, "Report title")
        workbook.write_subtitle(report, 1, 0, "Report subtitle")
        workbook.write_kpi(report, 2, 0, 3, "Indicators")
        create_excel_table(
            workbook,
            report,
            [{"value": "=SUM(A1:A2)", "count": 3}],
            start_row=5,
            start_col=0,
            columns=[ExcelColumn("value", "Value", 30), ExcelColumn("count", "Count", 12)],
        )
        write_meta_sheet(
            workbook,
            {
                "odysafe_threatmap_version": "0.1.0",
                "offline_mode": True,
                "command": "odysafe-threatmap report build report.txt",
            },
        )

    workbook = load_workbook(filename, data_only=False)
    report_sheet = workbook["Report"]
    assert workbook.sheetnames == ["Report", "_Meta"]
    assert workbook["_Meta"].sheet_state == "hidden"
    assert report_sheet.freeze_panes == "A7"
    assert report_sheet["A6"].value == "Value"
    assert report_sheet["A7"].value == "'=SUM(A1:A2)"
    assert report_sheet.tables
    assert report_sheet["A3"].value == "Indicators"
    manifest = json.loads((tmp_path / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["outputs"] == [str(filename)]
    assert manifest["attack_dataset"]["sha256"] == "def"


def test_summary_and_chart_are_written(tmp_path: Path) -> None:
    """Summary sheets contain KPI values and an inspectable bar chart."""
    filename = tmp_path / "summary.xlsx"
    with ExcelWorkbook(filename) as workbook:
        summary = write_summary_sheet(workbook, {"Techniques": 2}, [])
        create_bar_chart(
            workbook,
            summary,
            [{"Tactic": "Initial Access", "Count": 2}],
            start_row=6,
            start_col=0,
            title="Tactic coverage",
            x_label="Count",
            y_label="Tactic",
        )

    workbook = load_workbook(filename)
    summary_sheet = workbook["Synthèse"]
    assert summary_sheet["A3"].value == "Techniques"
    assert len(summary_sheet._charts) == 1
